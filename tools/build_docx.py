"""Build Word versions of the project's Markdown docs, with real equations.

These documents lean on notation that Markdown renders poorly -- mass-flow
dots, Greek letters, subscripted station numbers, and fractional exponents.
This converter emits **native Word equation objects (OMML)** for display
equations and genuine subscript/superscript runs for inline notation, rather
than Unicode lookalikes, so the result is editable as maths inside Word.

Run:
    python tools/build_docx.py            # all documents
    python tools/build_docx.py guide      # just the user guide
    python tools/build_docx.py validation

Requires python-docx. Verification (optional) additionally uses pywin32 to
drive Word and PyMuPDF to rasterise the result.
"""

from __future__ import annotations

import os
import re
import sys
from datetime import date

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)
WARN = RGBColor(0x9C, 0x27, 0x00)
CODE_BG = "F2F2F2"
CALLOUT_BG = "FFF8E1"

M_NS = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


# ---------------------------------------------------------------------------
# OMML construction
# ---------------------------------------------------------------------------


def _esc(t: str) -> str:
    return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def mr(t: str) -> str:
    """A maths run."""
    return f'<m:r><m:t xml:space="preserve">{_esc(t)}</m:t></m:r>'


def msub(base: str, sub: str) -> str:
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def msup(base: str, sup: str) -> str:
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"


def msubsup(base: str, sub: str, sup: str) -> str:
    return (
        f"<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub>"
        f"<m:sup>{sup}</m:sup></m:sSubSup>"
    )


def mfrac(num: str, den: str) -> str:
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def mrad(e: str) -> str:
    """Square root (degree hidden)."""
    return (
        '<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr>'
        f"<m:deg/><m:e>{e}</m:e></m:rad>"
    )


def macc(e: str, chr_: str = "&#775;") -> str:
    """Accent; default is a combining dot above, i.e. a rate."""
    return f'<m:acc><m:accPr><m:chr m:val="{chr_}"/></m:accPr><m:e>{e}</m:e></m:acc>'


def mdelim(e: str, beg: str = "(", end: str = ")") -> str:
    return (
        f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/>'
        f"</m:dPr><m:e>{e}</m:e></m:d>"
    )


DOT_M = macc(mr("m"))  # mass flow rate
DOT_R = macc(mr("r"))  # regression rate


#: Display equations, keyed by a marker planted in the Markdown scan.
EQUATIONS: dict[str, str] = {
    # SPI: mdot = Cd A sqrt(2 rho dP)
    "spi": (
        DOT_M + mr(" = ") + msub(mr("C"), mr("d")) + mr(" A ")
        + mrad(mr("2") + mr("ρ") + mr("Δ") + mr("P"))
    ),
    # Dyer non-equilibrium parameter
    "kappa": (
        mr("κ") + mr(" = ")
        + mrad(
            mfrac(
                msub(mr("P"), mr("1")) + mr(" − ") + msub(mr("P"), mr("2")),
                msub(mr("P"), mr("v")) + mr(" − ") + msub(mr("P"), mr("2")),
            )
        )
    ),
    # Dyer blend
    "dyer": (
        DOT_M + mr(" = ")
        + mfrac(mr("κ"), mr("1 + κ")) + msub(DOT_M, mr("SPI"))
        + mr(" + ")
        + mfrac(mr("1"), mr("1 + κ")) + msub(DOT_M, mr("HEM"))
    ),
    # Regression law
    "regression": (
        DOT_R + mr(" = a ") + msubsup(mr("G"), mr("ox"), mr("n"))
        + mr(" ") + msup(mr("L"), mr("m"))
    ),
    # Oxidiser mass flux
    "gox": (
        msub(mr("G"), mr("ox")) + mr(" = ")
        + mfrac(msub(DOT_M, mr("ox")), msub(mr("A"), mr("port")))
    ),
    # Ill-conditioned inversion
    "invert": (
        msub(DOT_M, mr("ox")) + mr(" ∝ ")
        + msup(mdelim(mr(" ⋯ "), "[", "]"), mfrac(mr("1"), mr("1 − n")))
    ),
    # Injector pressure drop
    "dp": (
        msub(mr("ΔP"), mr("inj")) + mr(" = ")
        + msub(mr("P"), mr("tank")) + mr(" − ") + msub(mr("P"), mr("chamber"))
    ),
    # Thermodynamic relation used to reconstruct entropy from the ESDU fit
    "tds": mr("T ds = dh − v dP"),
    # Discharge coefficient from a cold-flow measurement
    "cd": (
        msub(mr("C"), mr("d")) + mr(" = ")
        + mfrac(DOT_M, mr("A ") + mrad(mr("2") + mr("ρ") + mr("Δ") + mr("P")))
    ),
}


def add_display_equation(doc: Document, key: str, note: str = "") -> None:
    """Insert a centred display equation, optionally with a trailing note."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p._p.append(parse_xml(f"<m:oMathPara {M_NS}><m:oMath>{EQUATIONS[key]}</m:oMath></m:oMathPara>"))
    if note:
        np = doc.add_paragraph()
        np.alignment = WD_ALIGN_PARAGRAPH.CENTER
        np.paragraph_format.space_after = Pt(10)
        r = np.add_run(note)
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        r.italic = True


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------

# Variables that should render with a real subscript when they appear in prose.
INLINE_SUBS = [
    ("Ṁ_ox", "ṁ", "ox"),  # not used; placeholder for clarity
]

_SUB_PATTERN = re.compile(r"\b([A-Za-zΔṁṛ]{1,6})_([A-Za-z0-9]{1,8})\b")
_TOKEN = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+`|\[[^\]]+\]\([^)]+\))",
    re.DOTALL,
)


def _emit_with_subscripts(par, text: str, **fmt):
    """Add text to a paragraph, converting ``X_yyy`` into real subscript runs."""
    pos = 0
    for m in _SUB_PATTERN.finditer(text):
        if m.start() > pos:
            _run(par, text[pos:m.start()], **fmt)
        _run(par, m.group(1), **fmt)
        r = _run(par, m.group(2), **fmt)
        r.font.subscript = True
        pos = m.end()
    if pos < len(text):
        _run(par, text[pos:], **fmt)


def _run(par, text, bold=False, italic=False, code=False, size=None, color=None):
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    if code:
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xA3, 0x15, 0x15)
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


#: Characters that mark a backticked span as physics notation rather than a
#: code identifier. Names like ``mdot_ox`` or ``chamber_pressure`` contain none
#: of these, so they keep their monospace treatment; ``P₁ − P_sat`` does not.
_MATHY = set("ṁṙρκΔ∝√·−₁₂₀₃≈≥≤×⋯")


#: A lone variable, optionally subscripted, optionally with a value assigned:
#: ``a``, ``n``, ``L``, ``G_ox``, ``m = 0``, ``n = 0.681``. Deliberately narrow
#: so multi-word identifiers (``mdot_ox``, ``chamber_pressure``) stay code.
_VAR_SPAN = re.compile(r"[A-Za-zΔ](?:_[A-Za-z0-9]{1,4})?(?:\s*[=≈]\s*[-\d.]+)?$")


def _is_math_span(s: str) -> bool:
    return any(ch in _MATHY for ch in s) or bool(_VAR_SPAN.fullmatch(s.strip()))


def add_inline(par, text: str, base_bold=False, base_color=None, base_size=None):
    """Render Markdown inline markup into an existing paragraph."""
    for part in _TOKEN.split(text):
        if not part:
            continue
        fmt = dict(bold=base_bold, color=base_color, size=base_size)
        if part.startswith("**") and part.endswith("**"):
            _emit_with_subscripts(par, part[2:-2], **{**fmt, "bold": True})
        elif part.startswith("`") and part.endswith("`"):
            inner = part[1:-1]
            if _is_math_span(inner):
                # Physics notation: italic with real subscripts, not code.
                _emit_with_subscripts(par, inner, **{**fmt, "italic": True})
            else:
                _run(par, inner, code=True)
        elif part.startswith("[") and "](" in part:
            label = part[1:part.index("]")]
            _emit_with_subscripts(par, label, **{**fmt, "italic": True})
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            _emit_with_subscripts(par, part[1:-1], **{**fmt, "italic": True})
        else:
            _emit_with_subscripts(par, part, **fmt)


# ---------------------------------------------------------------------------
# Block helpers
# ---------------------------------------------------------------------------


def shade(el, fill: str):
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), fill)
    el.append(sh)


def add_code_block(doc: Document, lines: list[str]):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    shade(p._p.get_or_add_pPr(), CODE_BG)
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(9)


def add_callout(doc: Document, lines: list[str]):
    text = " ".join(ln.strip() for ln in lines).strip()
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.2)
    pf.right_indent = Inches(0.2)
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    shade(pPr, CALLOUT_BG)
    bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), "E0A800")
    bdr.append(left)
    pPr.append(bdr)
    add_inline(p, text, base_size=10)


def add_table(doc: Document, rows: list[list[str]]):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        shade(cell._tc.get_or_add_tcPr(), "DEEAF6")
        add_inline(cell.paragraphs[0], h, base_bold=True, base_size=9.5)
    for row in body:
        cells = t.add_row().cells
        for i, v in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], v, base_size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ---------------------------------------------------------------------------
# Front matter
# ---------------------------------------------------------------------------


def add_title_page(doc: Document, heading: str, subtitle: str, blurb: str):
    for _ in range(6):
        doc.add_paragraph()
    for line in ("N2O / Paraffin Hybrid", "Injector Sizing Tool"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.size = Pt(26)
        r.bold = True
        r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(16)
    r.font.color.rgb = GREY

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(blurb)
    r.italic = True
    r.font.size = Pt(11)
    r.font.color.rgb = GREY

    for _ in range(10):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"Generated {date.today():%d %B %Y}")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc: Document):
    p = doc.add_paragraph()
    r = p.add_run("Contents")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT

    p = doc.add_paragraph()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), r'TOC \o "1-3" \h \z \u')
    inner = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "Right-click here and choose “Update Field” to build the contents."
    inner.append(t)
    fld.append(inner)
    p._p.append(fld)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_numbers(doc: Document):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for instr in ("PAGE",):
        r = p.add_run()
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        r._r.append(fld)
        r2 = p.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r2._r.append(it)
        r3 = p.add_run()
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        r3._r.append(end)
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = GREY


# ---------------------------------------------------------------------------
# Markdown walk
# ---------------------------------------------------------------------------

#: Per-document configuration. ``quote_eq`` maps a substring of a blockquote
#: onto display equations; ``para_eq`` does the same for body paragraphs.
DOCS = {
    "guide": dict(
        src="USER_GUIDE.md",
        out="USER_GUIDE.docx",
        subtitle="User Guide",
        blurb="Background physics, operation, and interpretation of results",
        start_at=("# Part 1", "## Five-minute first run"),
        quote_eq=[
            ("ṁ = Cd", ["spi"]),
            ("κ = √", ["kappa", "dyer"]),
            # U+1E59 (r with dot ABOVE), the rate symbol, not U+1E5B (below).
            ("ṙ [mm/s]", ["regression", "gox"]),
        ],
        para_eq=[
            ("ṁ_ox ∝", ["invert"]),
            ("**ΔP_inj", ["dp"]),
        ],
    ),
    "validation": dict(
        src="VALIDATION.md",
        out="VALIDATION.docx",
        subtitle="Validation Report",
        blurb="What has been verified, what has not, and what it would take",
        start_at=None,
        quote_eq=[],
        # Note: `T ds = dh − v dP` is deliberately absent. It sits inside a
        # numbered list, where a display block would break the numbering, and
        # it reads perfectly well as inline italic maths.
        para_eq=[
            ("ṁ_ox ∝", ["invert"]),
            ("Cd = ṁ / (A√", ["cd"]),
        ],
    ),
}


def convert(src: str, doc: Document, cfg: dict):
    lines = src.split("\n")
    i = 0
    start_at = cfg.get("start_at")
    skipping = start_at is not None
    quote_eq = cfg.get("quote_eq", [])
    para_eq = cfg.get("para_eq", [])
    seen_eq: set[str] = set()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()

        # Drop any Markdown-only nav block before the real content starts.
        if skipping:
            if any(s.startswith(p) for p in start_at):
                skipping = False
            else:
                i += 1
                continue

        if not s:
            i += 1
            continue

        if s in ("---", "***", "___"):
            i += 1
            continue

        # Code block
        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, buf)
            continue

        # Table
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # Blockquote (equation or callout)
        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            joined = " ".join(buf)
            for trigger, keys in quote_eq:
                if trigger in joined:
                    for k in keys:
                        add_display_equation(doc, k)
                    break
            else:
                add_callout(doc, buf)
            continue

        # Headings
        if s.startswith("#"):
            level = len(s) - len(s.lstrip("#"))
            text = s.lstrip("#").strip()
            if level == 1:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=min(level, 3))
            h.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            add_inline(h, text)
            for r in h.runs:
                r.font.color.rgb = ACCENT
            i += 1
            continue

        # Bulleted list
        if re.match(r"^[-*]\s+", s):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                body = re.sub(r"^\s*[-*]\s+", "", lines[i])
                i += 1
                while (
                    i < len(lines)
                    and lines[i].startswith("  ")
                    and lines[i].strip()
                    # A new list item needs a marker *followed by space*; a
                    # continuation line may legitimately open with *emphasis*.
                    and not re.match(r"^\s*(?:[-*]\s|\d+\.\s)", lines[i])
                ):
                    body += " " + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, body)
            continue

        # Numbered list
        if re.match(r"^\d+\.\s+", s):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                body = re.sub(r"^\s*\d+\.\s+", "", lines[i])
                i += 1
                while (
                    i < len(lines)
                    and lines[i].startswith("   ")
                    and lines[i].strip()
                    and not re.match(r"^\s*\d+\.", lines[i].strip())
                ):
                    body += " " + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, body)
            continue

        # Paragraph (join wrapped lines)
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (
                not nxt
                or nxt.startswith(("#", ">", "|", "```", "---"))
                or re.match(r"^[-*]\s+|^\d+\.\s+", nxt)
            ):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        add_inline(p, " ".join(buf))

        # Promote selected inline expressions to display equations, once each.
        joined = " ".join(buf)
        for trigger, keys in para_eq:
            if trigger in joined:
                for k in keys:
                    if k not in seen_eq:
                        add_display_equation(doc, k)
                        seen_eq.add(k)
                break


def resolve(name: str) -> str:
    """Find a source document at the project root or one folder down.

    The Markdown sources have been reorganised into subfolders ("User Guide/",
    "Validation/"), so anchoring on the root alone breaks the build. Search
    both rather than hard-coding a layout that may move again.
    """
    direct = os.path.join(ROOT, name)
    if os.path.isfile(direct):
        return direct
    import glob

    hits = glob.glob(os.path.join(ROOT, "*", name))
    if hits:
        return hits[0]
    raise FileNotFoundError(
        f"could not find {name} at the project root or one level below it"
    )


def build(key: str) -> str:
    cfg = DOCS[key]
    doc = Document()

    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(7)

    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.9)
        sec.top_margin = sec.bottom_margin = Inches(0.85)

    add_title_page(doc, key, cfg["subtitle"], cfg["blurb"])
    add_toc(doc)
    add_page_numbers(doc)

    src_path = resolve(cfg["src"])
    with open(src_path, encoding="utf-8") as f:
        convert(f.read(), doc, cfg)

    # Write the Word file beside its Markdown source, wherever that lives.
    out = os.path.join(os.path.dirname(src_path), cfg["out"])
    doc.save(out)
    eqs = doc.element.xml.count("<m:oMathPara")
    print(f"wrote {os.path.relpath(out, ROOT)}  ({eqs} equations, {len(doc.tables)} tables)")
    return out


def main(argv: list[str]) -> int:
    keys = argv[1:] or list(DOCS)
    unknown = [k for k in keys if k not in DOCS]
    if unknown:
        print(f"unknown document(s): {', '.join(unknown)}", file=sys.stderr)
        print(f"choose from: {', '.join(DOCS)}", file=sys.stderr)
        return 2
    for k in keys:
        build(k)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
